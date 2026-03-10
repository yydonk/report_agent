from typing import Dict, Any, List
from docx import Document


def export_markdown(spec: Dict[str, Any]) -> str:
    f = spec.get("format_spec", {})
    c = spec.get("content_spec", {})
    lines: List[str] = []
    lines.append("# 模板规范清单")
    lines.append("## 格式规范")
    sections = f.get("sections", [])
    if sections:
        s0 = sections[0]
        lines.append(f"- 页面方向: {s0.get('orientation')}")
        m = s0.get("margins", {})
        lines.append(f"- 页边距: 上{m.get('top')} 下{m.get('bottom')} 左{m.get('left')} 右{m.get('right')}")
    ps = f.get("paragraph_styles", [])
    if ps:
        lines.append("- 段落样式:")
        for it in ps[:20]:
            lines.append(f"  - {it.get('style')}: 行距={it.get('line_spacing')} 段前={it.get('space_before')} 段后={it.get('space_after')}")
    ts = f.get("table_styles", [])
    if ts:
        lines.append("- 表格样式:")
        for it in ts:
            lines.append(f"  - {it.get('style')} ×{it.get('count')}")
    caps = f.get("captions", [])
    if caps:
        lines.append("- 图表标题样式:")
        for it in caps[:20]:
            lines.append(f"  - {it.get('style')} 编号={it.get('has_numbering')}")
    lines.append("## 内容规范")
    req = c.get("required_sections", [])
    if req:
        lines.append("- 必填章节: " + "、".join(req))
    dr = c.get("data_recording", [])
    if dr:
        lines.append("- 数据记录要求:")
        for t in dr[:20]:
            lines.append(f"  - {t}")
    am = c.get("analysis_methods", [])
    if am:
        lines.append("- 分析方法:")
        for t in am[:20]:
            lines.append(f"  - {t}")
    cr = c.get("conclusion_requirements", [])
    if cr:
        lines.append("- 结论要求:")
        for t in cr[:20]:
            lines.append(f"  - {t}")
    ph = c.get("placeholders", [])
    if ph:
        lines.append("- 占位与提示:")
        for t in ph[:20]:
            lines.append(f"  - {t}")
    return "\n".join(lines)


def export_docx(spec: Dict[str, Any], out_path: str):
    doc = Document()
    md = export_markdown(spec)
    for line in md.splitlines():
        doc.add_paragraph(line)
    doc.save(out_path)
